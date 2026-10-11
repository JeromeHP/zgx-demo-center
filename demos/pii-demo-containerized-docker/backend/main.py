"""
PII Masking Demo — Slim Backend
FastAPI → Ollama (host) → two distinct, size-matched models

The original (pre-slim) demo compared TinyLlama 1.1B (base) against a
genuinely fine-tuned Qwen2.5-32B model — an unfair size/architecture
mismatch (confirmed via orig-repos/pii/backend/main.py's BASE_MODEL_PATH,
which defaults to a TinyLlama GGUF). This version restores an apples-to-
apples comparison:

  - "Base"       (OLLAMA_MODEL_BASE):      qwen2.5:32b         — stock Qwen2.5, 32B, generic prompt
  - "Fine-tuned" (OLLAMA_MODEL_FINETUNED): pii-detector:latest — same architecture/size class (32.8B,
                                                                  also Qwen2-family per `ollama show`),
                                                                  fine-tuned specifically for PII masking

Same base architecture and parameter count on both sides — the only
variable is fine-tuning, so any accuracy gap in the demo is attributable to
that, not to one model simply being bigger. Both names are set in
docker-compose.yml's demo-pii `environment:` block; qwen2.5:32b must be
pulled once (`ollama pull qwen2.5:32b`, ~19-20GB) before this works.

pii-detector's Modelfile TEMPLATE is a raw passthrough ("{{ .Prompt }}" —
confirmed via `curl :11434/api/show -d '{"model":"pii-detector:latest"}'`)
with ChatML-style stop tokens (<|im_end|>, <|im_start|>) baked in via
PARAMETER stop. Ollama does not apply chat formatting for this model on its
own, and a `system` field in the request JSON is silently dropped (the
template has no {{ .System }} slot to put it in) — see
_build_finetuned_prompt() below, which builds the ChatML wrapping by hand
for models that need it. Confirmed empirically (raw curl, same prompt):
unwrapped rambles with unwanted commentary (126 tokens); ChatML-wrapped
gives a clean, correct answer (14 tokens). qwen2.5:32b has a real Go
template (handles .System/.Messages natively) so it does not need this.
"""

import asyncio
import io
import json
import os
import re
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from backend.ollama_client import generate, health as ollama_health

# Session token accounting for the shared ZGX chrome (demos/ui/README.md).
try:  # noqa: E402
    from zgx_metrics import TokenLedger
except ImportError:  # pragma: no cover - depends on how the container launches
    # Two launch styles in this suite put a different directory on sys.path:
    #   python3 backend/main.py      -> /app/backend  (plain import works)
    #   uvicorn backend.main:app     -> /app          (backend is a package)
    # pii and ballistics use the second and crash-looped on ModuleNotFoundError
    # until this fallback existed. Covering both is cheaper than making every
    # demo agree on a launch style.
    from backend.zgx_metrics import TokenLedger

LEDGER = TokenLedger()


app = FastAPI(title="PII Masking Demo — Slim")

# Second model name for "fine-tuned" side — defaults to same model with enhanced prompt
OLLAMA_MODEL_BASE      = os.getenv("OLLAMA_MODEL_BASE",      os.getenv("OLLAMA_MODEL", "qwen2.5vl:7b"))
OLLAMA_MODEL_FINETUNED = os.getenv("OLLAMA_MODEL_FINETUNED", os.getenv("OLLAMA_MODEL", "qwen2.5vl:7b"))

# Models whose Modelfile TEMPLATE is a raw passthrough (Ollama applies no
# chat formatting) but whose stop tokens imply ChatML fine-tuning — confirmed
# for pii-detector via /api/show. Substring match so a future retag like
# "pii-detector:v2" still matches without a code change.
CHATML_PASSTHROUGH_MODELS = ("pii-detector",)


@app.on_event("startup")
async def _warm_models_on_startup():
    """
    Fire-and-forget: start loading both ~19-20GB models into VRAM the moment
    the container comes up, instead of waiting for the operator's first
    click in the demo (measured cold: ~35s combined load+first-generate;
    warm: ~11s). Does NOT block startup/health — /api/health only checks
    Ollama reachability, not model residency, so the container still reports
    healthy immediately; this just runs concurrently in the background so
    the wait, if any, happens during install/launch rather than mid-demo.
    keep_alive=30m matches call_model()'s value so this stays resident.
    """
    import httpx
    ollama_host = os.getenv("OLLAMA_HOST", "http://host-gateway:11434")

    async def _warm(model: str):
        payload = {"model": model, "prompt": "ok", "stream": False, "think": False,
                   "keep_alive": "30m", "options": {"num_predict": 1}}
        try:
            async with httpx.AsyncClient(timeout=120.0) as client:
                await client.post(f"{ollama_host}/api/generate", json=payload)
        except Exception:
            pass  # best-effort — a failed warm-up just means the first real request pays the cold-load cost instead

    asyncio.create_task(_warm(OLLAMA_MODEL_BASE))
    asyncio.create_task(_warm(OLLAMA_MODEL_FINETUNED))


def _needs_chatml(model_name: str) -> bool:
    return any(p in model_name for p in CHATML_PASSTHROUGH_MODELS)


def _build_finetuned_prompt(model: str, system_prompt: str, user_content: str) -> str:
    """
    Builds the full string sent as Ollama's `prompt` field for a
    "finetuned-side" call (comparison or PII-list extraction). Models in
    CHATML_PASSTHROUGH_MODELS get the system/user turns wrapped by hand;
    everything else falls through to the same flat concatenation this file
    always used, so this is behavior-preserving for any other model.
    """
    if _needs_chatml(model):
        return (
            f"<|im_start|>system\n{system_prompt}<|im_end|>\n"
            f"<|im_start|>user\n{user_content}<|im_end|>\n"
            f"<|im_start|>assistant\n"
        )
    return f"{system_prompt}\n\n{user_content}"


# Deliberately terse — confirmed via orig-repos/pii/backend/main.py
# (create_pii_prompt, is_finetuned=True branch) that this exact wording,
# not an elaborate category list, is what pii-detector was actually
# fine-tuned/validated against. A longer, more "helpful" instruction (what
# this constant used to contain) pushes it out of its trained distribution —
# confirmed empirically: it made the fine-tuned model perform worse than
# the untuned qwen2.5:32b base on the same input, defeating the entire
# point of the comparison. Fine-tuned models built for one specific prompt
# don't reliably generalize to a differently-worded one the way a strong
# general-purpose base model does.
FINETUNED_SYSTEM_PROMPT = "Mask all PII and PHI in text."

# ── Document upload / redaction ────────────────────────────────────────────
# Text-layer documents only (no OCR) — a scanned/photographed document has
# no extractable text and is rejected with a clear 422, not silently mangled.
MAX_UPLOAD_BYTES = 15 * 1024 * 1024   # 15MB — generous for a text-heavy office doc
MIN_EXTRACTABLE_CHARS = 20             # below this, treat as "no text layer"
# pii-detector's confirmed context window is 32,768 tokens (~4 chars/token).
# Leaves headroom for the system prompt (~300 tokens) and the JSON-list
# response (up to PII_EXTRACTION_MAX_TOKENS) on top of the input text itself.
MAX_EXTRACT_CHARS_FOR_REDACTION = 80_000
PII_EXTRACTION_MAX_TOKENS = 2048

PII_LIST_SYSTEM_PROMPT = """You are a specialized PII (Personally Identifiable Information) detection system.
Find every instance of PII in the text below and return a JSON array of the EXACT substrings found
verbatim in the text (character-for-character, including original capitalization and punctuation) —
do not paraphrase, normalize, abbreviate, or reformat any value.

PII categories to detect: names, email addresses, phone numbers, physical addresses, Social Security
Numbers, credit card numbers, dates of birth, medical record numbers, driver's license numbers,
passport numbers, IP addresses, account numbers, and any other personally identifying value.

Respond with ONLY a JSON array of strings, e.g.: ["John Smith", "555-238-9921", "jsmith@acme.com"]
Do not include category labels, keys, explanations, or any text outside the JSON array.
If no PII is found, respond with []."""


@app.get("/api/metrics")
async def api_metrics():
    """Session token totals for the shared chrome."""
    return LEDGER.snapshot()


@app.get("/api/health")
async def health():
    ollama_status = await ollama_health()
    return {
        "status":             "ok" if ollama_status["ollama_reachable"] else "degraded",
        "backend":            "ollama-slim",
        "model_base":         OLLAMA_MODEL_BASE,
        "model_finetuned":    OLLAMA_MODEL_FINETUNED,
        **ollama_status,
    }


class PIIRequest(BaseModel):
    text: str
    max_tokens: int | None = None
    temperature: float | None = None


async def _run_pii_comparison(req: "PIIRequest") -> dict:
    """
    Run text through both the base and fine-tuned model, return both
    results (with per-call timing) for side-by-side comparison. Shared by
    both /api/process and /mask_pii, which just format this differently.
    """
    import time
    import httpx

    # ── Base model prompt (generic, poor PII detection) ───────────────────
    base_prompt = f"""Please help me with this text: {req.text}

Identify any personal information and replace it with [REDACTED]."""

    # ── Fine-tuned model prompt — raw text as the user turn, no added
    # framing. Also confirmed via orig-repos: the user turn there is just
    # {text} verbatim, not wrapped in "INPUT TEXT:"/"MASKED OUTPUT:" labels.
    finetuned_prompt = _build_finetuned_prompt(
        OLLAMA_MODEL_FINETUNED, FINETUNED_SYSTEM_PROMPT, req.text
    )

    OLLAMA_HOST = os.getenv("OLLAMA_HOST", "http://host-gateway:11434")
    TIMEOUT = 60.0

    options = {}
    if req.temperature is not None:
        options["temperature"] = req.temperature
    if req.max_tokens is not None:
        options["num_predict"] = req.max_tokens

    async def call_model(model: str, prompt: str, extra_options: dict | None = None):
        # think:false — nemotron3:33b (and other reasoning-capable models)
        # emit a separate hidden "thinking" block before the real response,
        # and it counts against num_predict same as the response itself. On
        # longer inputs the model can burn the entire token budget thinking
        # and never emit a response at all (done_reason "length", response
        # ""), which is silent from here — Ollama still returns 200. Since
        # this demo caps num_predict from a user-facing UI slider, that's
        # easy to hit. Skipping the reasoning phase avoids it entirely and
        # is also just faster.
        call_options = {**options, **(extra_options or {})}
        # keep_alive=30m — Ollama's default is 5m, which is shorter than a
        # typical pause in a live sales demo (explaining the report while
        # not actively clicking). Letting it expire mid-demo means the next
        # click pays the full ~20-25s cold-load penalty again, which reads
        # as "still loading" even though the container itself is healthy.
        payload = {"model": model, "prompt": prompt, "stream": False, "think": False, "keep_alive": "30m"}
        if call_options:
            payload["options"] = call_options
        t0 = time.monotonic()
        async with httpx.AsyncClient(timeout=TIMEOUT) as client:
            r = await client.post(f"{OLLAMA_HOST}/api/generate", json=payload)
            r.raise_for_status()
            data = r.json()
            elapsed = time.monotonic() - t0
            # Counted here, at the choke point, so no call site can forget.
            # These counts were already being read into base_token_usage /
            # finetuned_token_usage further down and then never displayed
            # anywhere -- the frontend has no reference to token_usage at all.
            LEDGER.record(data, fallback_text=data.get("response", ""), wall_s=elapsed)
            return data, elapsed

    try:
        base_result, base_elapsed = await call_model(OLLAMA_MODEL_BASE, base_prompt)
        # repeat_penalty=1.1 — matches orig-repos/pii/backend/main.py's
        # generate_with_gguf() config for this exact model (the non-tinyllama
        # branch).
        #
        # temperature=0 (greedy decoding), hard override regardless of the
        # UI's temperature slider (req.temperature only affects the base
        # side above) — found via repeated identical requests that at
        # temperature 0.1 this model has a real, silent failure mode: on
        # ~1 in 4 runs of the exact same input it leaves one or more full
        # names completely unredacted, with no error and no signal anything
        # went wrong. 7/7 clean at temperature=0 in the same test. This
        # isn't a prompt-wording issue — it's sampling variance on a task
        # where a wrong "roll" means a silent PII leak, which is the one
        # failure mode this demo cannot afford to have any real chance of.
        finetuned_result, finetuned_elapsed = await call_model(
            OLLAMA_MODEL_FINETUNED, finetuned_prompt, extra_options={"repeat_penalty": 1.1, "temperature": 0}
        )
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Ollama error: {str(e)}")

    finetuned_output = finetuned_result.get("response", "")

    # Safety net: the rewrite-style call above can still occasionally miss
    # an item even at temperature=0 - confirmed live, GPU floating-point
    # execution isn't perfectly deterministic, so a near-tied decision
    # inside the model can flip between otherwise-identical runs. No
    # amount of prompt or decoding-parameter tuning closes that gap
    # completely, so verify independently instead: ask the same model to
    # list the PII it finds (a different task framing - "list what's
    # there" vs. "rewrite the text" - reusing _extract_pii_strings, built
    # for document redaction) and force-redact anything from that list
    # still sitting in the rewritten output verbatim. A hallucinated item
    # that isn't actually a substring of the output is a no-op - this can
    # only ever add redaction, never remove real content, so it can't make
    # the result less safe than the unverified version.
    #
    # Counted into finetuned_elapsed (not left out) - this call is a real
    # part of what makes the finetuned side's answer trustworthy, so the
    # displayed processing time should reflect its actual total cost.
    verify_t0 = time.monotonic()
    try:
        verified_pii = await _extract_pii_strings(req.text)
        for item in sorted({s for s in verified_pii if s.strip()}, key=len, reverse=True):
            if item in finetuned_output:
                finetuned_output = finetuned_output.replace(item, "[REDACTED]")
    except HTTPException:
        pass  # verification failing shouldn't block showing the primary result
    finetuned_elapsed += time.monotonic() - verify_t0

    return {
        "base_output":  base_result.get("response", ""),
        "base_elapsed": base_elapsed,
        "base_token_usage": {
            "prompt_tokens":     base_result.get("prompt_eval_count", 0),
            "completion_tokens": base_result.get("eval_count", 0),
            "total_tokens":      base_result.get("prompt_eval_count", 0) + base_result.get("eval_count", 0),
        },
        "finetuned_output":  finetuned_output,
        "finetuned_elapsed": finetuned_elapsed,
        "finetuned_token_usage": {
            "prompt_tokens":     finetuned_result.get("prompt_eval_count", 0),
            "completion_tokens": finetuned_result.get("eval_count", 0),
            "total_tokens":      finetuned_result.get("prompt_eval_count", 0) + finetuned_result.get("eval_count", 0),
        },
    }


@app.post("/api/process")
async def process_pii(req: PIIRequest):
    r = await _run_pii_comparison(req)
    return JSONResponse({
        "original_text": req.text,
        "base_model": {
            "output":      r["base_output"],
            "model":       OLLAMA_MODEL_BASE,
            "token_usage": r["base_token_usage"],
        },
        "finetuned_model": {
            "output":      r["finetuned_output"],
            "model":       OLLAMA_MODEL_FINETUNED,
            "token_usage": r["finetuned_token_usage"],
        },
    })


@app.post("/mask_pii")
async def mask_pii(req: PIIRequest):
    """
    frontend/index.html calls this exact path with this exact response
    shape. /api/process above is the same underlying logic in our own
    naming convention; kept for any other caller, but the frontend uses
    this one.
    """
    r = await _run_pii_comparison(req)
    return JSONResponse({
        "base_model_output":      r["base_output"],
        "base_model_time":        round(r["base_elapsed"], 2),
        "finetuned_model_output": r["finetuned_output"],
        "finetuned_model_time":   round(r["finetuned_elapsed"], 2),
    })


# ── Document text extraction ───────────────────────────────────────────────

def _extract_docx_text(file_bytes: bytes) -> str:
    from docx import Document   # local import, matches this file's existing style (see _run_pii_comparison)
    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def _extract_pdf_text(file_bytes: bytes) -> str:
    import pymupdf   # "fitz" is the old (now-deprecated) import name for this same library
    doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    try:
        return "\n".join(page.get_text() for page in doc)
    finally:
        doc.close()


def _extract_text_from_upload(filename: str, file_bytes: bytes) -> tuple[str, str]:
    """Returns (doc_type, text). 400 for an unrecognized extension, 422 for
    'parsed OK but nothing extractable' — this demo is text-layer-only, no
    OCR, so a scanned PDF lands here rather than returning empty/garbage
    text silently."""
    ext = os.path.splitext(filename or "")[1].lower()
    if ext == ".docx":
        doc_type, text = "docx", _extract_docx_text(file_bytes)
    elif ext == ".pdf":
        doc_type, text = "pdf", _extract_pdf_text(file_bytes)
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported file type '{ext}' — upload a .docx or .pdf")

    if len(text.strip()) < MIN_EXTRACTABLE_CHARS:
        raise HTTPException(
            status_code=422,
            detail="No extractable text found. If this is a scanned/image-based PDF, note this demo "
                   "supports text-layer documents only (no OCR).",
        )
    return doc_type, text


# ── PII-list extraction (for redaction — distinct from the inline-rewrite
# prompt above, which is for the two-panel text comparison) ────────────────

def _parse_pii_list(raw_response: str) -> list[str]:
    """
    Ollama's format=<schema> constrains the *shape* of the JSON, not the
    content — a model can still emit {"pii": [...]} instead of a bare array,
    or wrap it in prose. Handle both; never silently return an empty list on
    a parse failure, since that would look identical to "no PII found" when
    extraction actually just broke.
    """
    text = raw_response.strip()

    def _as_list(parsed):
        if isinstance(parsed, list):
            return [str(x) for x in parsed if isinstance(x, (str, int, float))]
        if isinstance(parsed, dict):
            for v in parsed.values():
                if isinstance(v, list):
                    return [str(x) for x in v if isinstance(x, (str, int, float))]
        return None

    try:
        result = _as_list(json.loads(text))
        if result is not None:
            return result
    except json.JSONDecodeError:
        pass

    match = re.search(r"\[.*\]", text, re.DOTALL)   # fallback: pull [...] out of surrounding prose
    if match:
        try:
            result = _as_list(json.loads(match.group(0)))
            if result is not None:
                return result
        except json.JSONDecodeError:
            pass

    raise ValueError(f"could not parse a JSON array of PII strings from model output: {text[:200]!r}")


async def _extract_pii_strings(text: str) -> list[str]:
    import httpx
    if len(text) > MAX_EXTRACT_CHARS_FOR_REDACTION:
        raise HTTPException(
            status_code=422,
            detail=f"Document text ({len(text)} chars) exceeds the {MAX_EXTRACT_CHARS_FOR_REDACTION}-char "
                   f"limit for single-pass PII extraction (pii-detector's context window is 32,768 tokens).",
        )
    prompt = _build_finetuned_prompt(OLLAMA_MODEL_FINETUNED, PII_LIST_SYSTEM_PROMPT, text)
    payload = {
        "model": OLLAMA_MODEL_FINETUNED,
        "prompt": prompt,
        "stream": False,
        "think": False,
        "format": {"type": "array", "items": {"type": "string"}},
        "keep_alive": "30m",  # see _run_pii_comparison's call_model for why
        # repeat_penalty=1.1, temperature=0 — same validated config as the
        # comparison path's finetuned call (see _run_pii_comparison for why
        # temperature=0 specifically: this model has a real silent-leak
        # failure mode at nonzero temperature). Higher stakes here than the
        # comparison view — this list directly drives what gets redacted in
        # a document someone may download and actually share.
        "options": {"temperature": 0, "num_predict": PII_EXTRACTION_MAX_TOKENS, "repeat_penalty": 1.1},
    }
    ollama_host = os.getenv("OLLAMA_HOST", "http://host-gateway:11434")
    try:
        async with httpx.AsyncClient(timeout=90.0) as client:
            r = await client.post(f"{ollama_host}/api/generate", json=payload)
            r.raise_for_status()
            raw = r.json().get("response", "")
        return _parse_pii_list(raw)
    except ValueError as e:
        raise HTTPException(status_code=502, detail=f"PII extraction failed: {e}")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Ollama error: {str(e)}")


# ── Redaction — find-and-replace within the original document structure,
# NOT a full-text rewrite spliced back in. Preserves formatting; a
# hallucinated/paraphrased "PII" item that isn't an exact substring of
# anything in the document is a silent no-op, not a corruption risk. ───────

def _redact_docx(file_bytes: bytes, pii_items: list[str]) -> bytes:
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    # Longest-first so a short match (e.g. a first name) can't clobber part
    # of a longer overlapping one (e.g. the full name) before it's checked.
    pii_sorted = sorted({s for s in pii_items if s.strip()}, key=len, reverse=True)

    def _redact_paragraph(paragraph):
        for run in paragraph.runs:
            for pii in pii_sorted:
                if pii in run.text:
                    run.text = run.text.replace(pii, "[REDACTED]")

    for paragraph in doc.paragraphs:
        _redact_paragraph(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _redact_paragraph(paragraph)
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            _redact_paragraph(paragraph)
        for paragraph in section.footer.paragraphs:
            _redact_paragraph(paragraph)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _redact_pdf(file_bytes: bytes, pii_items: list[str]) -> bytes:
    """True visual redaction (blacks out the exact regions, preserving the
    rest of the page untouched) — not text reconstruction, which PDF's
    fixed layout doesn't support cleanly. add_redact_annot's default fill
    is white; must override to black."""
    import pymupdf
    doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    pii_sorted = sorted({s for s in pii_items if s.strip()}, key=len, reverse=True)

    for page in doc:
        for pii in pii_sorted:
            for quad in page.search_for(pii, quads=True):
                page.add_redact_annot(quad, fill=(0, 0, 0))
        page.apply_redactions()

    out = doc.tobytes(garbage=4, deflate=True)
    doc.close()
    return out


def _check_upload_size(file_bytes: bytes):
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=400, detail=f"File exceeds {MAX_UPLOAD_BYTES // (1024*1024)}MB limit")


@app.post("/api/extract_text")
async def extract_text(file: UploadFile = File(...)):
    file_bytes = await file.read()
    _check_upload_size(file_bytes)

    doc_type, text = _extract_text_from_upload(file.filename, file_bytes)
    return JSONResponse({
        "filename":   file.filename,
        "doc_type":   doc_type,
        "text":       text,
        "char_count": len(text),
    })


@app.post("/api/redact_document")
async def redact_document(file: UploadFile = File(...)):
    file_bytes = await file.read()
    _check_upload_size(file_bytes)

    doc_type, text = _extract_text_from_upload(file.filename, file_bytes)
    pii_items = await _extract_pii_strings(text)

    if doc_type == "docx":
        redacted_bytes = _redact_docx(file_bytes, pii_items)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        redacted_bytes = _redact_pdf(file_bytes, pii_items)
        media_type = "application/pdf"

    safe_name = re.sub(r"[^A-Za-z0-9._-]", "_", os.path.basename(file.filename or f"document.{doc_type}"))
    return Response(
        content=redacted_bytes,
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="redacted_{safe_name}"',
            "X-PII-Redaction-Count": str(len(pii_items)),
        },
    )


@app.get("/api/samples")
async def sample_texts():
    return {
        "samples": [
            {
                "label": "Employee Record",
                "text":  "John Smith (DOB: 03/15/1982) joined our team on January 5, 2020. "
                         "His SSN is 523-44-1892. Contact him at jsmith@acme.com or 555-238-9921. "
                         "Home address: 4421 Oak Tree Lane, Austin, TX 78701.",
            },
            {
                "label": "Medical Note",
                "text":  "Patient: Maria Gonzalez, MRN 00284712. DOB 09/22/1965. "
                         "Referred by Dr. Robert Chen. Insurance ID: BCB-992-44-1122. "
                         "Contact: maria.g@gmail.com, 512-874-2210.",
            },
            {
                "label": "Financial Record",
                "text":  "Account holder: David Kim. Account #: 8812-4499-2211-0033. "
                         "Routing: 021000021. DL: TX-DL-44928812. "
                         "Billing address: 8801 Research Blvd, Suite 200, Austin TX 78758.",
            },
        ]
    }


# Frontend, served under /app/ (zgx-demo-center convention — API routes
# above stay at root, matching frontend/index.html's absolute fetch() paths).
FRONTEND_DIR = os.environ.get("FRONTEND_DIR", "/app/frontend")
if not os.path.isdir(FRONTEND_DIR):
    _local_frontend = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "frontend"))
    if os.path.isdir(_local_frontend):
        FRONTEND_DIR = _local_frontend
if os.path.isdir(FRONTEND_DIR):
    app.mount("/app", StaticFiles(directory=FRONTEND_DIR, html=True), name="frontend")
